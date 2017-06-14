# -*- coding: utf-8 -*-

import json

from django.test import TestCase

from openpyxl import load_workbook
from StringIO import StringIO
from string import ascii_uppercase

from seqr.models import Family, Individual
from seqr.views.utils.export_table_utils import export_table, export_families, export_individuals, _parse_phenotips_data


class ExportTableUtilsTest(TestCase):
    fixtures = ['users', '1kg_project']

    def test_export_table(self):
        header = ['column1', 'column2']
        rows = [['row1_v1', 'row1_v2'], ['row2_v1', 'row2_v2']]

        # test tsv format
        response = export_table('test_file', header, rows, file_format='tsv')

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content, '\n'.join(['\t'.join(row) for row in [header]+rows]) + '\n')

        # test Excel format
        response = export_table('test_file', header, rows, file_format='xls')
        self.assertEqual(response.status_code, 200)
        wb = load_workbook(StringIO(response.content))
        worksheet = wb.active

        self.assertListEqual([cell.value for cell in worksheet['A']], ['Column1', 'row1_v1', 'row2_v1'])
        self.assertListEqual([cell.value for cell in worksheet['B']], ['Column2', 'row1_v2', 'row2_v2'])
        self.assertEqual([cell.value for cell in worksheet['C']], [None, None, None])

        # test unknown format
        self.assertRaisesRegexp(ValueError, '.*format.*',
            lambda: export_table('test_file', header, rows, file_format='unknown_format')
        )

    def test_export_families(self):
        test_families = Family.objects.all()

        # test tsv with all columns
        response = export_families('test_families_table', test_families, 'tsv', include_project_column=True, include_case_review_columns=True)
        self.assertEqual(response.status_code, 200)
        rows = [row.split('\t') for row in response.content.rstrip('\n').split('\n')]
        HEADER = ['project', 'family_id', 'display_name', 'created_date', 'description', 'analysis_status',  'analysis_summary', 'analysis_notes', 'internal_case_review_summary', 'internal_case_review_notes']
        self.assertListEqual(rows[0], HEADER)
        self.assertEqual(len(rows), 13)

        # test tsv without project column
        response = export_families('test_families_table', test_families, 'tsv', include_project_column=False, include_case_review_columns=True)
        self.assertEqual(response.status_code, 200)
        rows = [row.split('\t') for row in response.content.rstrip('\n').split('\n')]
        HEADER = ['family_id', 'display_name', 'created_date', 'description', 'analysis_status', 'analysis_summary', 'analysis_notes', 'internal_case_review_summary', 'internal_case_review_notes']
        self.assertListEqual(rows[0], HEADER)
        self.assertEqual(len(rows), 13)

        # test tsv without case review columns
        response = export_families('test_families_table', test_families, 'tsv', include_project_column=False, include_case_review_columns=False)
        self.assertEqual(response.status_code, 200)
        rows = [row.split('\t') for row in response.content.rstrip('\n').split('\n')]
        HEADER = ['family_id', 'display_name', 'created_date', 'description', 'analysis_status', 'analysis_summary', 'analysis_notes']
        self.assertListEqual(rows[0], HEADER)
        self.assertEqual(len(rows), 13)

        # test Excel format
        response = export_families('test_families_table', test_families, 'xls', include_project_column=True, include_case_review_columns=True)
        self.assertEqual(response.status_code, 200)
        wb = load_workbook(StringIO(response.content))
        worksheet = wb.active
        self.assertListEqual([cell.value for cell in worksheet['A']], ['Project'] + [u'1kg project nåme with uniçøde']*12)
        self.assertListEqual([cell.value for cell in worksheet['B']], ['Family Id', '1', '2', '3', '4', '5', '6', '7', '8', '9', '10', '11', '12'])
        self.assertListEqual([cell.value for cell in worksheet['C']], ['Display Name', '1', '2', '3', '4', '5', '6', '7', '8', '9', '10', '11', '12'])

        # test unknown format
        self.assertRaisesRegexp(ValueError, '.*format.*',
            lambda: export_families('test_families_table', test_families, file_format='unknown_format')
        )

    def test_export_individuals(self):
        test_individuals = Individual.objects.filter(individual_id='NA19675')

        # test tsv with all columns
        response = export_individuals(
            'test_individuals_table',
            test_individuals,
            'tsv',
            include_project_column=True,
            include_case_review_columns=True,
            include_phenotips_columns=True,
        )
        self.assertEqual(response.status_code, 200)
        rows = [row.split('\t') for row in response.content.rstrip('\n').split('\n')]
        HEADER = [
            'project', 'family_id', 'individual_id', 'paternal_id', 'maternal_id', 'sex', 'affected_status', 'notes',
            'case_review_status', 'case_review_status_last_modified_date', 'case_review_status_last_modified_by', 'case_review_discussion',
            'phenotips_features_present', 'phenotips_features_absent', 'paternal_ancestry', 'maternal_ancestry', 'age_of_onset'
        ]

        self.assertListEqual(rows[0], HEADER)
        self.assertEqual(len(rows), 2)

        # test Excel format
        response = export_individuals('test_families_table', test_individuals, 'xls', include_project_column=True, include_case_review_columns=True, include_phenotips_columns=True)
        self.assertEqual(response.status_code, 200)
        wb = load_workbook(StringIO(response.content))
        worksheet = wb.active

        # test unknown format
        self.assertRaisesRegexp(ValueError, '.*format.*',
            lambda: export_individuals('test_families_table', test_individuals, file_format='unknown_format')
        )

    def test_parse_phenotips_data(self):
        test_individuals = Individual.objects.filter(individual_id='NA19675')

        phenotips_json = json.loads(test_individuals[0].phenotips_data)

        parsed_data = _parse_phenotips_data(phenotips_json)

        self.assertSetEqual(
            set(parsed_data.keys()),
            {'age_of_onset', 'candidate_genes', 'maternal_ancestry', 'paternal_ancestry', 'phenotips_features_absent', 'phenotips_features_present', 'previously_tested_genes'}
        )

        self.assertEqual(parsed_data['age_of_onset'], 'Adult onset')
        self.assertEqual(parsed_data['candidate_genes'], 'EHBP1L1 (comments EHBP1L1), ACY3 (comments ACY3)')
        self.assertEqual(parsed_data['paternal_ancestry'], 'African Americans')
        self.assertEqual(parsed_data['maternal_ancestry'], 'Azerbaijanis')
        self.assertEqual(parsed_data['phenotips_features_absent'], 'Arrhythmia, Complete atrioventricular canal defect, Failure to thrive')
        self.assertEqual(parsed_data['phenotips_features_present'], 'Defect in the atrial septum, Morphological abnormality of the central nervous system, Tetralogy of Fallot')
        self.assertEqual(parsed_data['previously_tested_genes'], 'IKBKAP (comments IKBKAP), CCDC102B (comments for CCDC102B)')
